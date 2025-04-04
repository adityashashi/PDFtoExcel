import os
import fitz  # PyMuPDF
import pdfplumber
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Side

def extract_tables_from_pdf(pdf_path, output_excel_path):
    wb = Workbook()
    ws_index = 0

    with pdfplumber.open(pdf_path) as pdf:
        for page_number, page in enumerate(pdf.pages, start=1):
            tables = page.extract_tables()

            if not tables or len(tables) == 0:
                # fallback to PyMuPDF text-based table detection
                doc = fitz.open(pdf_path)
                page_fitz = doc[page_number - 1]
                words = page_fitz.get_text("words")  # list of (x0, y0, x1, y1, word, block_no, line_no, word_no)
                words.sort(key=lambda w: (round(w[1], 1), w[0]))  # sort by y, then x

                # group words into lines by y coordinate
                lines = {}
                for w in words:
                    y = round(w[1], 1)
                    lines.setdefault(y, []).append(w)

                sorted_lines = sorted(lines.items())
                table_data = []
                for y, words_line in sorted_lines:
                    line_text = [w[4] for w in sorted(words_line, key=lambda x: x[0])]
                    table_data.append(line_text)

                sheet = wb.create_sheet(title=f"Page_{page_number}_text")
                for r_idx, row in enumerate(table_data, start=1):
                    for c_idx, cell in enumerate(row, start=1):
                        sheet.cell(row=r_idx, column=c_idx, value=cell)
                        sheet.cell(row=r_idx, column=c_idx).alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
                continue

            for table in tables:
                # normalize the table structure to retain layout
                cleaned_table = [[cell if cell is not None else "" for cell in row] for row in table]
                max_columns = max(len(row) for row in cleaned_table)
                formatted_table = [row + [''] * (max_columns - len(row)) for row in cleaned_table]

                sheet = wb.create_sheet(title=f"Page_{page_number}_table{ws_index}")
                ws_index += 1
                for r_idx, row in enumerate(formatted_table, start=1):
                    for c_idx, cell in enumerate(row, start=1):
                        sheet.cell(row=r_idx, column=c_idx, value=cell)
                        sheet.cell(row=r_idx, column=c_idx).alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
                        sheet.cell(row=r_idx, column=c_idx).border = Border(
                            left=Side(style='thin'),
                            right=Side(style='thin'),
                            top=Side(style='thin'),
                            bottom=Side(style='thin')
                        )

    if "Sheet" in wb.sheetnames:
        std = wb["Sheet"]
        wb.remove(std)

    wb.save(output_excel_path)

# Example usage:
pdf_path = "test2.pdf"  # Replace with your PDF file
output_excel = "output.xlsx"
extract_tables_from_pdf(pdf_path, output_excel)
print(f"Tables extracted and saved to {output_excel}")